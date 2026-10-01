"""
Builds the Microsoft Ignite 2026 deck for the equinix-arc-demo repository.

    python docs/presentation/build_deck.py [output.pptx]

Requires python-pptx and qrcode (pip install python-pptx qrcode[pil]).
Content mirrors docs/PRESENTATION.md (slide outline + speaker notes) - edit both together.

Design: midnight title/closing slides, light content slides, Equinix-red reserved for
"the private path", a red left-edge bar on content slides, numbered route-marker circles
for sequences. Fonts: Bahnschrift SemiBold (titles), Segoe UI (body), Cascadia Code (code).
"""
import io
import sys
from pathlib import Path

import qrcode
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# --- Palette ------------------------------------------------------------------------
NIGHT = "0B1622"      # dominant dark
NIGHT2 = "15263A"     # dark cards
PAPER = "F4F6F9"      # light background
CARD = "FFFFFF"
INK = "16202C"        # dark text
MUTED = "5E6B7A"
RULE = "D6DDE5"
LIGHT = "C9D3DD"      # light text on dark
SKY = "6CC7E0"        # secondary accent (Azure-ish, on dark)
RED = "E4002B"        # the private path / Equinix
AZURE = "0078D4"
AWS = "FF9900"
GREEN = "1E9E6A"
AMBER = "C77700"
CODE_BG = "0F1B2A"
CODE_FG = "D7E3EE"

TITLE_FONT = "Bahnschrift SemiBold"
BODY_FONT = "Segoe UI"
CODE_FONT = "Cascadia Code"
SYMBOL_FONT = "Segoe UI Symbol"

W, H = 13.333, 7.5
REPO_URL = "https://github.com/bbenz/equinix-arc-demo"
FOOTER = "Microsoft Ignite 2026  \u00b7  One fleet, three footprints"


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


# --- Primitives -------------------------------------------------------------------
def shape(slide, kind, x, y, w, h, fill=None, line=None, line_w=1.0, radius=None):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_w)
    else:
        s.line.fill.background()
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0):
    return shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, fill, line, line_w)


def rounded(slide, x, y, w, h, fill=None, line=None, line_w=1.0, radius=0.08):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, line, line_w, radius)


def circle(slide, cx, cy, r, fill=None, line=None, line_w=1.5):
    return shape(slide, MSO_SHAPE.OVAL, cx - r, cy - r, 2 * r, 2 * r, fill, line, line_w)


def connector(slide, x1, y1, x2, y2, color, width=2.0, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return c


def _style_run(run, size, color, font, bold, italic):
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.font.name = font
    run.font.bold = bold
    run.font.italic = italic


def text(slide, x, y, w, h, paragraphs, size=16, color=INK, font=BODY_FONT, bold=False, italic=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.0, space_after=0, target=None):
    """paragraphs: str | list of str | list of list[(text, {opts})]. opts: size,color,font,bold,italic."""
    if target is None:
        target = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = target.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = anchor
    if isinstance(paragraphs, str):
        paragraphs = [paragraphs]
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        runs = para if isinstance(para, list) else [(para, {})]
        for run_text, opts in runs:
            r = p.add_run()
            r.text = run_text
            _style_run(r, opts.get("size", size), opts.get("color", color), opts.get("font", font),
                       opts.get("bold", bold), opts.get("italic", italic))
    return target


def bullets(slide, x, y, w, h, items, size=15, color=INK, bullet_color=RED, space_after=8):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space_after)
        marker = p.add_run()
        marker.text = "\u25A0  "
        _style_run(marker, size - 5, bullet_color, BODY_FONT, False, False)
        r = p.add_run()
        r.text = item
        _style_run(r, size, color, BODY_FONT, False, False)
    return tb


def notes(slide, text_value):
    slide.notes_slide.notes_text_frame.text = text_value


def light_slide(prs, title, number):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(PAPER)
    rect(s, 0, 0, 0.14, H, fill=RED)  # motif: red left edge = "the private path"
    text(s, 0.75, 0.42, 11.8, 0.8, title, size=34, color=INK, font=TITLE_FONT, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 0.75, 6.85, 8.0, 0.3, FOOTER, size=10, color=MUTED)
    text(s, 11.6, 6.85, 1.0, 0.3, str(number), size=10, color=MUTED, align=PP_ALIGN.RIGHT)
    return s


def dark_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(NIGHT)
    return s


def marker(slide, cx, cy, label, r=0.3, fill=NIGHT, color="FFFFFF", line=RED, size=14, font=TITLE_FONT):
    c = circle(slide, cx, cy, r, fill=fill, line=line, line_w=2.0)
    text(slide, 0, 0, 0, 0, [[(label, {})]], size=size, color=color, font=font, bold=False,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, target=c)
    return c


def code_card(slide, x, y, w, h, caption, code, size=11.5):
    rect(slide, x, y, w, h, fill=CODE_BG)
    text(slide, x + 0.25, y + 0.15, w - 0.5, 0.3, caption, size=10, color=SKY, font=CODE_FONT)
    text(slide, x + 0.25, y + 0.5, w - 0.5, h - 0.65, code.strip("\n").split("\n"), size=size,
         color=CODE_FG, font=CODE_FONT)


def card(slide, x, y, w, h, band=None, fill=CARD):
    rect(slide, x, y, w, h, fill=fill, line=RULE, line_w=0.75)
    if band:
        rect(slide, x, y, w, 0.1, fill=band)


def table(slide, x, y, col_widths, rows, header_fill=NIGHT, size=12, row_h=0.46, first_col_bold=True,
          highlight_col=None):
    gf = slide.shapes.add_table(len(rows), len(col_widths), Inches(x), Inches(y),
                                Inches(sum(col_widths)), Inches(row_h * len(rows)))
    tbl = gf.table
    tbl.horz_banding = False
    for i, cw in enumerate(col_widths):
        tbl.columns[i].width = Inches(cw)
    for r_i, row in enumerate(rows):
        tbl.rows[r_i].height = Inches(row_h)
        for c_i, value in enumerate(row):
            cell = tbl.cell(r_i, c_i)
            cell.fill.solid()
            if r_i == 0:
                cell.fill.fore_color.rgb = rgb(header_fill)
            elif highlight_col is not None and c_i == highlight_col:
                cell.fill.fore_color.rgb = rgb("FDF0F3")
            else:
                cell.fill.fore_color.rgb = rgb(CARD if r_i % 2 else "EEF2F6")
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = value
            is_header = r_i == 0
            _style_run(r, size, "FFFFFF" if is_header else INK, BODY_FONT,
                       is_header or (first_col_bold and c_i == 0), False)
    return tbl


def qr_png(url):
    img = qrcode.make(url, border=1)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return buf


# --- Slides -------------------------------------------------------------------------
def slide_title(prs):
    s = dark_slide(prs)
    text(s, 0.75, 1.15, 6.9, 0.4, "MICROSOFT IGNITE 2026  \u00b7  SAN FRANCISCO", size=13, color=SKY, bold=True)
    text(s, 0.75, 1.6, 6.9, 1.9, ["One fleet,", "three footprints"], size=54, color="FFFFFF", font=TITLE_FONT)
    text(s, 0.75, 3.6, 6.6, 1.0, ["AKS, EKS and Kubernetes at Equinix \u2014", "private over ExpressRoute"],
         size=22, color=LIGHT)
    text(s, 0.75, 4.85, 6.6, 0.7, ["Azure Kubernetes Fleet Manager  \u00b7  Azure Arc",
                                   "Equinix Fabric  \u00b7  Azure ExpressRoute"], size=13, color="8FA1B3")
    text(s, 0.75, 6.25, 6.6, 0.4, "Brian Benz  \u00b7  Microsoft", size=15, color="FFFFFF")

    hub = (10.05, 3.55)
    aks, eks, eqx = (8.15, 1.55), (11.95, 1.55), (10.05, 6.05)
    connector(s, *hub, *aks, SKY, 2.5)
    connector(s, *hub, *eks, "8FA1B3", 2.5, dash=True)
    connector(s, *hub, *eqx, RED, 5)
    circle(s, *hub, 0.82, fill=NIGHT2, line=SKY, line_w=2)
    text(s, hub[0] - 0.8, hub[1] - 0.4, 1.6, 0.8, ["Fleet", "Manager"], size=15, color="FFFFFF", font=TITLE_FONT,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    for (cx, cy), fill, label, fg in [(aks, AZURE, "AKS", "FFFFFF"), (eks, AWS, "EKS", NIGHT), (eqx, RED, "Equinix", "FFFFFF")]:
        circle(s, cx, cy, 0.58, fill=fill)
        text(s, cx - 0.6, cy - 0.3, 1.2, 0.6, label, size=15 if label != "Equinix" else 13, color=fg, font=TITLE_FONT,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 7.55, 2.45, 1.6, 0.3, "native member", size=11, color=SKY)
    text(s, 10.95, 2.45, 1.55, 0.3, "Arc \u00b7 internet", size=11, color="8FA1B3", align=PP_ALIGN.RIGHT)
    text(s, 10.45, 4.95, 2.3, 0.55, ["Arc \u00b7 ExpressRoute", "no internet egress"], size=11, color=RED, bold=True)
    notes(s, "Same app, three very different places, one control plane - and the Equinix footprint never "
             "touches the internet. Over the next 20 minutes: why, how it's wired, and a live demo.")


def slide_sprawl(prs):
    s = light_slide(prs, "Kubernetes everywhere \u2014 and the sprawl", 2)
    cols = [
        (AZURE, "AZ", "Azure \u00b7 AKS", ["Cloud-native apps", "AI and data services", "Native Fleet member"]),
        (AWS, "AWS", "AWS \u00b7 EKS", ["Acquired or existing estates", "Team and contract choices", "Different LB, IAM, tooling"]),
        (RED, "EQX", "Colocation \u00b7 Equinix", ["Data gravity and latency", "Sovereignty, GPUs, partners", "Often no internet egress"]),
    ]
    for i, (color, tag, heading, lines) in enumerate(cols):
        x = 0.75 + i * 4.1
        card(s, x, 1.55, 3.75, 3.1, band=color)
        marker(s, x + 0.65, 2.25, tag, r=0.36, fill=color, line=color, color="FFFFFF" if color != AWS else NIGHT, size=12)
        text(s, x + 1.2, 2.0, 2.45, 0.5, heading, size=19, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        bullets(s, x + 0.35, 3.0, 3.15, 1.6, lines, size=15, bullet_color=color, space_after=10)
    rect(s, 0.75, 5.1, 11.95, 1.1, fill=NIGHT)
    text(s, 1.1, 5.1, 11.3, 1.1, [[("Result:  ", {"color": LIGHT}), ("N", {"color": RED, "bold": True}),
                                    (" control planes  \u00b7  ", {}), ("N", {"color": RED, "bold": True}),
                                    (" pipelines  \u00b7  ", {}), ("N", {"color": RED, "bold": True}),
                                    (" kinds of drift", {})]],
         size=22, color="FFFFFF", font=TITLE_FONT, anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "Most enterprises run Kubernetes in more than one place. Colocation is common for data gravity, "
             "latency, sovereignty and GPU capacity - and those sites often forbid internet egress by policy.")


def slide_goals(prs):
    s = light_slide(prs, "What good looks like", 3)
    goals = [
        ("1", "One control plane", "Azure Kubernetes Fleet Manager hub for every cluster, wherever it runs."),
        ("2", "One app definition", "Cloud-neutral manifests, placed everywhere, never forked."),
        ("3", "Differences as data", "Label-selected overrides for load balancers, banners and settings."),
        ("4", "Private where it matters", "Colocation traffic to Azure never touches the public internet."),
    ]
    for i, (num, head, body) in enumerate(goals):
        x = 0.75 + (i % 2) * 6.1
        y = 1.6 + (i // 2) * 2.5
        accent = RED if num == "4" else NIGHT
        if num == "4":
            rect(s, x, y, 5.85, 2.2, fill="FDE9EE", line=RED, line_w=1.0)
        else:
            card(s, x, y, 5.85, 2.2)
        marker(s, x + 0.75, y + 1.1, num, r=0.42, fill=accent, line=accent, size=20)
        text(s, x + 1.45, y + 0.42, 4.2, 0.5, head, size=21, bold=True, color=INK)
        text(s, x + 1.45, y + 0.97, 4.2, 0.95, body, size=15, color=MUTED if num != "4" else INK)
    notes(s, "Four goals drive everything that follows. The fourth one - private where it matters - is what "
             "Equinix Fabric and ExpressRoute add to the Fleet + Arc story.")


def slide_architecture(prs):
    s = light_slide(prs, "Architecture", 4)
    rect(s, 4.4, 1.45, 4.5, 1.05, fill=NIGHT)
    text(s, 4.4, 1.5, 4.5, 0.5, "Fleet Manager hub", size=19, color="FFFFFF", font=TITLE_FONT, align=PP_ALIGN.CENTER)
    text(s, 4.4, 1.98, 4.5, 0.4, "ClusterResourcePlacement  \u00b7  ResourceOverrides", size=12, color=LIGHT, align=PP_ALIGN.CENTER)
    members = [
        (0.85, AZURE, "aks-demo", ["Azure \u00b7 westus2", "Native Fleet member"], AZURE, False),
        (4.85, AWS, "eks-demo", ["AWS \u00b7 us-west-2", "Arc over the internet"], "8FA1B3", True),
        (8.85, RED, "equinix-demo", ["K3s in an Equinix SV cage", "Arc + Arc gateway over ExpressRoute"], RED, False),
    ]
    for x, band, name, lines, line_color, dashed in members:
        connector(s, 6.65, 2.5, x + 1.8, 3.5, line_color, 4.5 if band == RED else 2.25, dash=dashed)
    for x, band, name, lines, line_color, dashed in members:
        card(s, x, 3.6, 3.6, 1.55, band=band)
        text(s, x + 0.3, 3.8, 3.0, 0.45, name, size=18, bold=True, font=CODE_FONT)
        text(s, x + 0.3, 4.3, 3.1, 0.8, lines, size=13, color=MUTED)
    rounded(s, 8.6, 5.4, 4.1, 1.15, fill="FDE9EE", line=RED, line_w=1.25, radius=0.06)
    text(s, 8.8, 5.4, 3.75, 1.15, [[("No internet egress. ", {"bold": True, "color": RED})],
                                  "Equinix Fabric \u2192 ExpressRoute \u2192 hub VNet \u2192 egress proxy \u2192 Arc gateway"],
         size=12, color=INK, anchor=MSO_ANCHOR.MIDDLE, space_after=2)
    text(s, 0.85, 5.4, 7.4, 1.15, ["Every member runs its own full copy of Online Boutique, with its own Redis.",
                                   "The hub places it; no request ever crosses clusters."], size=14, color=MUTED, space_after=4,
         anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "Two ways to reach Azure: the internet for EKS, a private circuit for Equinix. Fleet doesn't care "
             "which - both are Arc-enabled members. Each storefront is fully independent.")


def slide_private_path(prs):
    s = dark_slide(prs)
    text(s, 0.75, 0.45, 11.8, 0.8, "The private path", size=36, color="FFFFFF", font=TITLE_FONT, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 0.75, 1.2, 11.8, 0.5, "Zero internet egress from the cage. End-to-end TLS \u2014 no inspection.", size=17, color=LIGHT)
    centers = [1.3 + i * 1.75 for i in range(7)]
    cy = 3.35
    pad = 0.75
    rounded(s, centers[0] - pad, 2.3, centers[2] - centers[0] + 2 * pad, 2.75, fill=NIGHT2, radius=0.04)
    rounded(s, centers[4] - pad, 2.3, centers[6] - centers[4] + 2 * pad, 2.75, fill=NIGHT2, radius=0.04)
    text(s, centers[0] - 0.65, 2.4, 3.6, 0.35, "EQUINIX SV CAGE", size=11, color=RED, bold=True)
    text(s, centers[4] - 0.65, 2.4, 3.6, 0.35, "AZURE HUB  \u00b7  WESTUS2", size=11, color=SKY, bold=True)
    text(s, centers[3] - 0.9, 2.4, 1.8, 0.35, "MICROSOFT EDGE", size=10, color=LIGHT, bold=True, align=PP_ALIGN.CENTER)
    connector(s, centers[0], cy, centers[6], cy, RED, 4)
    hops = [
        ("K3s node", "Arc + Fleet agents, containerd"),
        ("Edge router", "or Fabric Cloud Router \u00b7 BGP"),
        ("Equinix Fabric", "redundant virtual connections"),
        ("ExpressRoute", "private peering \u00b7 VLAN 200"),
        ("ER gateway", "hub VNet 10.50.0.0/16"),
        ("Egress proxy", "Squid \u00b7 passthrough \u00b7 allowlist"),
        ("Arc gateway", "*.gw.arc.azure.com"),
    ]
    for i, (head, desc) in enumerate(hops):
        marker(s, centers[i], cy, str(i + 1), r=0.36, fill=NIGHT, line=RED, size=16)
        text(s, centers[i] - 0.82, 3.9, 1.64, 0.4, head, size=13, color="FFFFFF", bold=True, align=PP_ALIGN.CENTER)
        text(s, centers[i] - 0.82, 4.3, 1.64, 0.95, desc, size=11, color=LIGHT, align=PP_ALIGN.CENTER)
    text(s, 0.75, 5.55, 11.8, 0.9, [[("Proof on stage:  ", {"color": RED, "bold": True}),
                                     ("BGP routes learned by Azure  \u00b7  Arc gateway enabled  \u00b7  proxy log entries from Equinix node IPs", {})]],
         size=15, color="FFFFFF")
    notes(s, "No internet egress in the cage. Arc agents use an explicit passthrough proxy that is only reachable "
             "over ExpressRoute, and they only need the nine Arc gateway endpoints. TLS stays end to end.")


def slide_as_code(prs):
    s = light_slide(prs, "Equinix Fabric + ExpressRoute, as code", 6)
    code_card(s, 0.75, 1.45, 6.25, 2.45, "terraform/azure/expressroute.tf", '''
resource "azurerm_express_route_circuit" "equinix" {
  service_provider_name = "Equinix"
  peering_location      = "Silicon Valley"  # SV1
  bandwidth_in_mbps     = 50
  sku {
    tier   = "Standard"
    family = "MeteredData"
  }
}''')
    code_card(s, 0.75, 4.1, 6.25, 2.45, "terraform/equinix/main.tf", '''
resource "equinix_fabric_connection" "azure_primary" {
  type = "EVPL_VC"        # IP_VC from a Cloud Router
  z_side { access_point {
    type               = "SP"
    authentication_key = var.expressroute_service_key
    peering_type       = "PRIVATE"
    link_protocol { type = "QINQ"  vlan_c_tag = 200 }
  } }
}''')
    items = [
        ("1", NIGHT, "Three origin types", "Fabric port \u00b7 Fabric Cloud Router \u00b7 service token"),
        ("2", NIGHT, "Same as Quick Connect", "Portal: Microsoft Azure \u2192 Azure ExpressRoute \u2192 paste service key"),
        ("3", RED, "Mind the meter", "ExpressRoute bills from the moment the service key is issued"),
        ("4", NIGHT, "Secrets stay secret", "Service key passed in-process only; Equinix API keys never on disk"),
    ]
    for i, (num, color, head, body) in enumerate(items):
        y = 1.55 + i * 1.27
        marker(s, 7.75, y + 0.4, num, r=0.3, fill=color, line=color, size=14)
        text(s, 8.3, y, 4.3, 0.42, head, size=17, bold=True, color=RED if color == RED else INK)
        text(s, 8.3, y + 0.5, 4.3, 0.7, body, size=13, color=MUTED)
    notes(s, "Both sides are Terraform: azurerm for the circuit, gateway and peering; the Equinix provider for the "
             "Fabric connections. Quick Connect in the portal does the same thing. Remember the circuit bills from "
             "the moment its service key is issued.")


def slide_arc_gateway(prs):
    s = light_slide(prs, "Why the Arc gateway?", 7)
    rect(s, 0.75, 1.55, 3.7, 4.95, fill=NIGHT)
    text(s, 0.75, 1.8, 3.7, 1.6, "9", size=110, color="FFFFFF", font=TITLE_FONT, align=PP_ALIGN.CENTER)
    text(s, 1.05, 3.65, 3.1, 1.0, "FQDNs to allow through the proxy \u2014 instead of dozens", size=16, color=LIGHT,
         align=PP_ALIGN.CENTER)
    text(s, 1.05, 4.95, 3.1, 1.2, "Required by Fleet Manager for Arc members behind a proxy", size=13, color=SKY,
         align=PP_ALIGN.CENTER, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    options = [
        (GREEN, "\u2713", "Arc gateway + passthrough proxy over ExpressRoute",
         "Supported path for proxied Fleet members \u00b7 one auditable egress point \u00b7 no internet in the cage"),
        ("5F6B78", "\u2715", "Arc Private Link for Kubernetes",
         "Still preview \u00b7 Entra ID, ARM and MCR still over the internet \u00b7 no Cluster Connect"),
        ("5F6B78", "\u2715", "ExpressRoute Microsoft peering",
         "Public NAT prefixes and route filters \u00b7 doesn't cover every Arc endpoint"),
    ]
    for i, (color, glyph, head, body) in enumerate(options):
        y = 1.55 + i * 1.75
        card(s, 4.85, y, 7.85, 1.5, fill=CARD if i else "F0FAF5")
        c = circle(s, 5.5, y + 0.75, 0.36, fill=color)
        text(s, 0, 0, 0, 0, glyph, size=20, color="FFFFFF", font=SYMBOL_FONT, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE, target=c)
        text(s, 6.15, y + 0.2, 6.35, 0.45, head, size=17, bold=True)
        text(s, 6.15, y + 0.68, 6.35, 0.75, body, size=13, color=MUTED)
    notes(s, "The most common architecture question, answered up front. Arc Private Link for Kubernetes is still "
             "preview and still needs Entra ID, ARM and MCR over the internet. Fleet requires the Arc gateway for "
             "Arc members behind a proxy - and the proxy must be passthrough, not TLS-terminating.")


def slide_overrides(prs):
    s = light_slide(prs, "One definition. Differences as data.", 8)
    code_card(s, 0.75, 1.5, 5.0, 3.85, "kubernetes/fleet/cluster-resource-placement.yaml", '''
kind: ClusterResourcePlacement
metadata:
  name: crp-online-boutique
spec:
  resourceSelectors:
    - kind: Namespace
      name: online-boutique
  policy:
    placementType: PickAll
    # every member labeled
    # demo=equinix-arc-online-boutique''', size=12)
    rows = [
        ["Override", "Azure", "AWS", "Equinix"],
        ["LB Service", "health probe", "NLB, internet-facing", "node IPs over ER"],
        ["ENV_PLATFORM", "azure", "aws", "onprem"],
        ["Store banner", "Azure", "AWS", "On-Premises"],
        ["Redis marker", "cloud=azure", "cloud=aws", "cloud=equinix"],
    ]
    table(s, 6.15, 1.5, [1.65, 1.45, 1.85, 1.6], rows, size=12.5, row_h=0.62, highlight_col=3)
    text(s, 6.15, 4.75, 6.55, 1.0, ["Selected by member labels: cloud \u00b7 provider \u00b7 connectivity \u00b7 site.",
                                    "A fourth site = one more rule per override."], size=14, color=MUTED, space_after=4)
    rounded(s, 0.75, 5.7, 11.95, 0.75, fill="FDE9EE", line=RED, line_w=1, radius=0.1)
    text(s, 1.05, 5.7, 11.35, 0.75, [[("The base never forks. ", {"bold": True, "color": RED}),
                                       ("kubernetes/base has zero cloud references \u2014 Fleet applies the differences at placement time.", {})]],
         size=14, anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "One placement selects every member of the demo. Three small overrides express everything that differs: "
             "load balancer annotations, the platform banner (Azure, AWS, On-Premises) and a Redis marker.")


def slide_demo(prs):
    s = dark_slide(prs)
    text(s, 0.75, 0.9, 6, 1.3, "Demo", size=66, color="FFFFFF", font=TITLE_FONT)
    text(s, 0.75, 2.15, 9, 0.5, "About 12 minutes \u00b7 everything shown is in the repo", size=17, color=LIGHT)
    steps = ["Fleet members and labels", "Proof of the private path", "One definition + overrides",
             "Three storefronts", "Live change via Arc Cluster Connect"]
    centers = [1.5 + i * 2.45 for i in range(5)]
    connector(s, centers[0], 4.35, centers[-1], 4.35, RED, 4)
    for i, label in enumerate(steps):
        marker(s, centers[i], 4.35, str(i + 1), r=0.42, fill=NIGHT, line=RED, size=20)
        text(s, centers[i] - 1.1, 5.0, 2.2, 0.9, label, size=15, color="FFFFFF", align=PP_ALIGN.CENTER)
    text(s, 0.75, 6.45, 11.8, 0.45, [[("Fallbacks ready:  ", {"color": RED, "bold": True}),
                                      ("pre-captured proof panel  \u00b7  Arc Cluster Connect port-forward  \u00b7  recorded video", {})]],
         size=13, color=LIGHT)
    notes(s, "Follow docs/DEMO-RUNSHEET.md. Fallbacks: pre-captured show-private-path output, Arc Cluster Connect "
             "port-forward for the Equinix storefront, and the recorded backup video.")


def slide_proof(prs):
    s = light_slide(prs, "What you just saw", 10)
    proofs = [
        ("Fabric: PROVISIONED", "Equinix Fabric \u00b7 redundant connections"),
        ("BGP: routes learned", "10.80.0.0/24 at the Azure ER gateway"),
        ("Arc: Connected", "through the Arc gateway"),
        ("Proxy log", "Arc tunnels from Equinix node IPs"),
        ("Storefront", "fetched from Azure over ExpressRoute"),
        ("kubectl via Arc", "no VPN, no inbound ports"),
    ]
    for i, (head, body) in enumerate(proofs):
        x = 0.75 + (i % 3) * 4.05
        y = 1.6 + (i // 3) * 2.2
        card(s, x, y, 3.8, 1.95)
        c = circle(s, x + 0.6, y + 0.72, 0.33, fill=GREEN)
        text(s, 0, 0, 0, 0, "\u2713", size=18, color="FFFFFF", font=SYMBOL_FONT, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE, target=c)
        text(s, x + 1.1, y + 0.42, 2.6, 0.62, head, size=16, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.35, y + 1.2, 3.3, 0.6, body, size=13, color=MUTED)
    text(s, 0.75, 6.1, 11.9, 0.45, "Captured by scripts/demo-show-private-path.ps1 and scripts/09-validate-demo.ps1",
         size=12, color=MUTED, font=CODE_FONT)
    notes(s, "Recap each proof point: Fabric status, BGP routes, Arc gateway, proxy log, the storefront fetched across "
             "the circuit, and kubectl through Arc Cluster Connect with no VPN.")


def slide_good_to_know(prs):
    s = light_slide(prs, "Good to know", 11)
    columns = [
        (GREEN, "Works today", [
            "Placement + overrides on Arc members (GA)",
            "Arc gateway for Arc-enabled Kubernetes",
            "Fabric \u2194 ExpressRoute: portal, API, Terraform",
            "Arc Cluster Connect for day-2 access",
        ]),
        (AMBER, "Plan around", [
            "Fleet update runs are AKS-only",
            "Passthrough proxies only \u2014 no TLS termination",
            "Equinix Metal EOL (June 2026): use colocation",
            "AKS on bare metal: preview, single node, East US",
        ]),
    ]
    for i, (color, head, items) in enumerate(columns):
        x = 0.75 + i * 6.1
        card(s, x, 1.55, 5.85, 3.55, band=color)
        text(s, x + 0.4, 1.85, 5.0, 0.55, head, size=22, bold=True, color=color, font=TITLE_FONT)
        bullets(s, x + 0.4, 2.65, 5.2, 2.3, items, size=15, bullet_color=color, space_after=16)
    text(s, 0.75, 5.5, 11.9, 0.45, "Details: docs/ARCHITECTURE.md  \u00b7  docs/EQUINIX-CLUSTER.md  \u00b7  docs/NETWORKING-EXPRESSROUTE.md",
         size=12, color=MUTED, font=CODE_FONT)
    notes(s, "Be upfront about the limits - it builds trust. Placement and overrides are GA for Arc members; update "
             "runs are AKS-only. Equinix Metal is gone, so this demo runs on colocated hardware.")


def slide_cost(prs):
    s = light_slide(prs, "Cost, production, and the code", 12)
    stats = [("$0.89/hr", ["Azure + AWS run rate", "list price, excludes Equinix"], INK),
             ("$55/mo", ["50 Mbps ExpressRoute circuit", "bills from creation"], INK),
             ("0", ["internet egress paths", "from the Equinix cage"], RED)]
    for i, (big, label, color) in enumerate(stats):
        x = 0.75 + i * 4.05
        card(s, x, 1.5, 3.8, 1.9)
        text(s, x + 0.3, 1.58, 3.2, 1.0, big, size=44, color=color, font=TITLE_FONT, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.3, 2.6, 3.35, 0.7, label, size=13, color=MUTED)
    text(s, 0.75, 3.75, 7.5, 0.5, "Production checklist", size=19, bold=True, font=TITLE_FONT)
    bullets(s, 0.75, 4.35, 7.6, 2.3, [
        "Azure Firewall explicit proxy, or an HA proxy pair",
        "ExpressRoute Metro or two peering locations",
        "BGP MD5 + BFD, summarized prefixes",
        "Private Fleet hub (Arc gateway already in place)",
        "Monitor and Defender via Arc extensions",
    ], size=15, space_after=7)
    s.shapes.add_picture(qr_png(REPO_URL), Inches(9.55), Inches(3.7), Inches(2.0), Inches(2.0))
    text(s, 8.55, 5.85, 4.0, 0.45, "github.com/bbenz/equinix-arc-demo", size=13, bold=True, align=PP_ALIGN.CENTER,
         font=CODE_FONT)
    notes(s, "About $0.89 an hour for Azure and AWS plus $55 a month for the circuit; the Equinix side is quoted "
             "separately. Call to action: clone the repo and try rehearsal mode today - no Equinix hardware needed.")


def slide_thanks(prs):
    s = dark_slide(prs)
    text(s, 0.75, 0.9, 8, 1.2, "Thank you", size=60, color="FFFFFF", font=TITLE_FONT)
    text(s, 0.75, 2.05, 11.8, 0.5, "Fleet for consistency  \u00b7  Arc for reach  \u00b7  Equinix + ExpressRoute for a private path",
         size=17, color=LIGHT)
    links = [
        ("Sample repo", "github.com/bbenz/equinix-arc-demo"),
        ("Fleet Manager + Arc", "learn.microsoft.com/azure/kubernetes-fleet/concepts-fleet-arc-integration"),
        ("Arc gateway", "learn.microsoft.com/azure/azure-arc/kubernetes/arc-gateway-simplify-networking"),
        ("Equinix \u2194 ExpressRoute", "docs.equinix.com/fabric-marketplace/connecting-to-service-provider/microsoft/azure-qc"),
    ]
    for i, (label, url) in enumerate(links):
        y = 3.05 + i * 0.85
        marker(s, 1.05, y + 0.27, str(i + 1), r=0.27, fill=NIGHT, line=RED, size=13)
        text(s, 1.55, y, 2.6, 0.55, label, size=16, color="FFFFFF", bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text(s, 4.15, y, 8.6, 0.55, url, size=11.5, color="9AD3E3", font=CODE_FONT, anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "Close: one control plane, three footprints, private where it matters - and all of it is code you can "
             "run today.")


def slide_appendix_labels(prs):
    s = light_slide(prs, "Appendix \u00b7 labels and egress allowlist", 14)
    rows = [
        ["Label", "aks-demo", "eks-demo", "equinix-demo"],
        ["cloud", "azure", "aws", "equinix"],
        ["provider", "aks", "eks", "k3s"],
        ["connectivity", "azure-native", "public-internet", "expressroute"],
        ["site", "azure-westus2", "aws-us-west-2", "equinix-sv"],
        ["demo", "equinix-arc-online-boutique (all members)", "", ""],
    ]
    tbl = table(s, 0.75, 1.5, [1.55, 1.6, 1.75, 1.7], rows, size=12, row_h=0.55, highlight_col=3)
    tbl.cell(5, 1).merge(tbl.cell(5, 3))
    card(s, 7.75, 1.5, 4.95, 4.75)
    text(s, 8.05, 1.7, 4.4, 0.45, "Egress proxy allowlist", size=17, bold=True)
    groups = [
        ("Arc gateway", ".gw.arc.azure.com \u00b7 management.azure.com \u00b7 .obo.arc.azure.com \u00b7 login.microsoftonline.com \u00b7 .login.microsoft.com \u00b7 .his.arc.azure.com \u00b7 mcr.microsoft.com \u00b7 .data.mcr.microsoft.com"),
        ("Fleet hub", ".azmk8s.io"),
        ("Images", "us-central1-docker.pkg.dev \u00b7 Docker Hub (registry, auth, CDN)"),
        ("K3s install", "get.k3s.io \u00b7 update.k3s.io \u00b7 github.com"),
    ]
    paragraphs = []
    for head, body in groups:
        paragraphs.append([(head, {"bold": True, "color": RED, "size": 13})])
        paragraphs.append([(body, {"size": 11, "font": CODE_FONT, "color": INK})])
    text(s, 8.05, 2.3, 4.45, 3.8, paragraphs, size=12, space_after=6)
    notes(s, "Reference: the label schema used by placement and overrides, and the full egress allowlist on the "
             "Azure proxy. Everything else is denied.")


def slide_appendix_pipeline(prs):
    s = light_slide(prs, "Appendix \u00b7 the pipeline", 15)
    steps = [("00", "check-tools"), ("00b", "bootstrap-auth"), ("01", "test-access"), ("02", "select-regions"),
             ("03", "plan"), ("04", "apply  $"), ("05", "connect-er  $"), ("06", "connect-arc"),
             ("07", "join-fleet"), ("08", "deploy-workload"), ("09", "validate"), ("99", "destroy")]
    for row in range(2):
        y = 2.3 + row * 2.2
        xs = [1.4 + i * 2.05 for i in range(6)]
        connector(s, xs[0], y, xs[-1], y, RED, 3)
        for i in range(6):
            num, label = steps[row * 6 + i]
            billable = "$" in label
            marker(s, xs[i], y, num, r=0.42, fill=RED if billable else NIGHT, line=RED, size=16)
            text(s, xs[i] - 1.0, y + 0.62, 2.0, 0.5, "make " + label.replace("  $", ""), size=11, color=INK,
                 align=PP_ALIGN.CENTER, font=CODE_FONT)
    text(s, 0.75, 6.05, 11.9, 0.45, [[("Red = billable step.    ", {"color": RED, "bold": True}),
                                      ("Every step is idempotent; 99 waits for Equinix to release the circuit before deleting it.", {})]],
         size=13, color=MUTED)
    notes(s, "The repo's numbered pipeline. Only apply and connect-er create billable resources, and both ask for "
             "typed confirmation.")


def build(output):
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "One fleet, three footprints: AKS, EKS and Kubernetes at Equinix"
    prs.core_properties.author = "Brian Benz"
    prs.core_properties.subject = "Microsoft Ignite 2026 - Azure Kubernetes Fleet Manager + Azure Arc + Equinix Fabric + ExpressRoute"
    for fn in (slide_title, slide_sprawl, slide_goals, slide_architecture, slide_private_path, slide_as_code,
               slide_arc_gateway, slide_overrides, slide_demo, slide_proof, slide_good_to_know, slide_cost,
               slide_thanks, slide_appendix_labels, slide_appendix_pipeline):
        fn(prs)
    prs.save(output)
    print(f"wrote {output} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    default = Path(__file__).with_name("equinix-arc-fleet-ignite-2026.pptx")
    build(sys.argv[1] if len(sys.argv) > 1 else str(default))
